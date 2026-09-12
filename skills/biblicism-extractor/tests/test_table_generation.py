from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import docx


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = SKILL_ROOT / "scripts"
WORD_MAKER_PATH = SCRIPTS_DIR / "word_maker.py"
GENERATE_OUTPUT_PATH = SCRIPTS_DIR / "generate_output.py"
RUN_PIPELINE_PATH = SCRIPTS_DIR / "run_pipeline.py"


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load module from {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def build_template(path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Original Title")

    table = document.add_table(rows=2, cols=5)
    headers = [
        "фрагмент текста",
        "библеизм",
        "структурная классификация",
        "особенности функционирования",
        "источник и пояснение",
    ]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header

    template_row = table.rows[1]
    for index, cell in enumerate(template_row.cells):
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(f"TEMPLATE-{index}")
        run.bold = True
        run.italic = True

    second_table = document.add_table(rows=1, cols=1)
    second_table.rows[0].cells[0].text = "obsolete"
    document.save(str(path))


def write_data(path: Path) -> None:
    payload = [
        {
            "fragment": "He carried his cross.",
            "biblicism": "carry one's cross",
            "classification": "библейские фразеологизмы",
            "features": "устойчивое выражение",
            "source_explanation": "Евангельский образ несения креста.",
        },
        {
            "fragment": "Miriam spoke softly.",
            "biblicism": "Miriam",
            "classification": "лексические единицы (антропонимы)",
            "features": "ветхозаветное имя",
            "source_explanation": "Имя пророчицы Мариам.",
        },
    ]
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


class TableGenerationTests(unittest.TestCase):
    def test_word_maker_preserves_template_table_structure(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            template = temp / "example.docx"
            output = temp / "result.docx"
            data = temp / "data.json"
            source = temp / "american-gods.pdf"
            source.write_text("dummy", encoding="utf-8")

            build_template(template)
            write_data(data)

            result = subprocess.run(
                [
                    sys.executable,
                    "-X",
                    "utf8",
                    str(WORD_MAKER_PATH),
                    str(template),
                    str(output),
                    str(data),
                    str(source),
                ],
                capture_output=True,
                text=True,
                encoding="utf-8",
                check=False,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertTrue(output.exists())

            generated = docx.Document(str(output))
            self.assertEqual(len(generated.tables), 1)
            paragraph_texts = [
                paragraph.text
                for paragraph in generated.paragraphs
                if paragraph.text.strip()
            ]
            self.assertEqual(paragraph_texts, ["Original Title"])

            table = generated.tables[0]
            self.assertEqual(len(table.columns), 5)
            self.assertEqual(len(table.rows), 3)
            self.assertEqual(table.rows[1].cells[0].text, "He carried his cross.")
            self.assertEqual(table.rows[2].cells[1].text, "Miriam")

            self.assertEqual(table.rows[0].cells[0].text, "фрагмент текста")
            self.assertEqual(table.rows[0].cells[4].text, "источник и пояснение")
            self.assertNotIn(
                "obsolete",
                "\n".join(cell.text for row in table.rows for cell in row.cells),
            )

    def test_generate_output_fails_instead_of_writing_fake_docx(self) -> None:
        module = load_module(GENERATE_OUTPUT_PATH, "generate_output_test")

        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            template = temp / "example.docx"
            output = temp / "result.docx"
            data = temp / "data.json"
            source = temp / "american-gods.pdf"
            source.write_text("dummy", encoding="utf-8")

            build_template(template)
            write_data(data)

            backup_json = output.with_suffix(".json")
            backup_csv = output.with_suffix(".csv")
            backup_md = output.with_suffix(".md")

            fake_result = subprocess.CompletedProcess(
                args=[], returncode=0, stdout="", stderr=""
            )

            with patch.object(module, "run_word_maker", return_value=fake_result):
                with patch.object(
                    sys,
                    "argv",
                    [
                        str(GENERATE_OUTPUT_PATH),
                        str(template),
                        str(output),
                        str(data),
                        str(source),
                    ],
                ):
                    with self.assertRaises(SystemExit) as error:
                        module.main()

            self.assertEqual(error.exception.code, 1)
            self.assertFalse(output.exists())
            self.assertTrue(backup_json.exists())
            self.assertTrue(backup_csv.exists())
            self.assertTrue(backup_md.exists())

    def test_run_pipeline_finalize_creates_template_based_docx(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = Path(temp_dir) / "workspace"
            source_dir = workspace / "source"
            example_dir = workspace / "example"
            temp_out_dir = workspace / "temp"
            output_dir = workspace / "output"

            source_dir.mkdir(parents=True)
            example_dir.mkdir(parents=True)
            temp_out_dir.mkdir(parents=True)
            output_dir.mkdir(parents=True)

            template = example_dir / "example.docx"
            data = temp_out_dir / "data.json"
            source = source_dir / "american-gods.pdf"
            source.write_text("dummy", encoding="utf-8")

            build_template(template)
            write_data(data)

            result = subprocess.run(
                [
                    sys.executable,
                    "-X",
                    "utf8",
                    str(RUN_PIPELINE_PATH),
                    "finalize",
                    "--workspace",
                    str(workspace),
                ],
                capture_output=True,
                text=True,
                encoding="utf-8",
                check=False,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

            final_docx = output_dir / "american-gods.docx"
            self.assertTrue(final_docx.exists())

            generated = docx.Document(str(final_docx))
            self.assertEqual(len(generated.tables), 1)
            self.assertEqual(len(generated.tables[0].columns), 5)
            self.assertEqual(len(generated.tables[0].rows), 3)
            self.assertEqual(
                generated.tables[0].rows[1].cells[1].text, "carry one's cross"
            )
            self.assertTrue((output_dir / "american-gods.json").exists())
            self.assertTrue((output_dir / "american-gods.csv").exists())
            self.assertTrue((output_dir / "american-gods.md").exists())


if __name__ == "__main__":
    unittest.main()
