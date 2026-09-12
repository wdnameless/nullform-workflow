from __future__ import annotations

import sys
from pathlib import Path


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: verify_artifact.py <file-path>")

    path = Path(sys.argv[1])
    if not path.exists():
        raise SystemExit(f"MISSING: {path}")

    if path.is_dir():
        print(f"DIR_OK: {path}")
        return

    size = path.stat().st_size
    print(f"FILE_OK: {path} SIZE={size}")

    if path.suffix.lower() == ".docx":
        import docx

        doc = docx.Document(str(path))
        print(f"DOCX_TABLES={len(doc.tables)} PARAGRAPHS={len(doc.paragraphs)}")
        if doc.tables:
            table = doc.tables[0]
            print(f"FIRST_TABLE_ROWS={len(table.rows)} COLS={len(table.columns)}")
    elif path.suffix.lower() in {".txt", ".md", ".json", ".csv"}:
        with path.open("rb") as handle:
            first_bytes = handle.read(16)
        print("FIRST_BYTES=" + " ".join(f"{byte:02X}" for byte in first_bytes))


if __name__ == "__main__":
    main()
