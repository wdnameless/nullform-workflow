from __future__ import annotations

import json
import shutil
import sys
from copy import deepcopy
from pathlib import Path
from typing import Any


def normalize_text(value: object) -> str:
    return " ".join(str(value or "").split())


def normalize_rows(data: object) -> list[list[str]]:
    rows: list[list[str]] = []
    seen: set[tuple[str, ...]] = set()

    if not isinstance(data, list):
        return rows

    for item in data:
        if isinstance(item, dict):
            row = [
                normalize_text(item.get("fragment", "")),
                normalize_text(item.get("biblicism", "")),
                normalize_text(item.get("classification", "")),
                normalize_text(item.get("features", "")),
                normalize_text(item.get("source_explanation", "")),
            ]
        elif isinstance(item, list):
            row = [normalize_text(cell) for cell in item[:5]]
            while len(row) < 5:
                row.append("")
        else:
            continue

        key = tuple(row)
        if key in seen or not any(row):
            continue
        seen.add(key)
        rows.append(row)

    return rows


def clear_cell(cell: Any) -> None:
    cell.text = ""


def clone_row(table: Any, row_template: Any) -> Any:
    new_tr = deepcopy(row_template)
    table._tbl.append(new_tr)
    return table.rows[len(table.rows) - 1]


def clear_table_data_rows(table: Any) -> None:
    for index in range(len(table.rows) - 1, 0, -1):
        table._tbl.remove(table.rows[index]._tr)


def cleanup_template_body(doc: Any) -> None:
    body = doc._body._element
    children = list(body)
    sect_pr = None
    for child in children:
        if child.tag.endswith("sectPr"):
            sect_pr = child
            break

    keep_first_paragraph = True
    keep_first_table = True
    for child in children:
        tag = child.tag.split("}")[-1]
        if tag == "sectPr":
            continue
        if tag == "p" and keep_first_paragraph:
            keep_first_paragraph = False
            continue
        if tag == "tbl" and keep_first_table:
            keep_first_table = False
            continue
        body.remove(child)

    if sect_pr is not None and body[-1] is not sect_pr:
        body.remove(sect_pr)
        body.append(sect_pr)


def write_rows_into_first_table(doc: Any, rows: list[list[str]]) -> None:
    if not doc.tables:
        raise SystemExit("Template does not contain tables.")

    table = doc.tables[0]
    if len(table.rows) < 2:
        raise SystemExit("Template must contain a header row and one sample data row.")

    row_template = deepcopy(table.rows[1]._tr)
    clear_table_data_rows(table)

    for row_data in rows:
        row = clone_row(table, row_template)
        for cell in row.cells:
            clear_cell(cell)
        for index, value in enumerate(row_data[: len(row.cells)]):
            row.cells[index].text = value


def main() -> None:
    if len(sys.argv) not in {4, 5}:
        raise SystemExit(
            "Usage: word_maker.py <template.docx> <output.docx> <data.json> [source-file]"
        )

    template_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    data_path = Path(sys.argv[3])

    shutil.copy2(template_path, output_path)

    import docx

    doc = docx.Document(str(output_path))
    cleanup_template_body(doc)
    rows = normalize_rows(json.loads(data_path.read_text(encoding="utf-8")))
    write_rows_into_first_table(doc, rows)
    doc.save(str(output_path))


if __name__ == "__main__":
    main()
