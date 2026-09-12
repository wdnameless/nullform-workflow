from __future__ import annotations

import csv
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side


HEADERS = [
    "фрагмент текста",
    "библеизм",
    "структурная классификация",
    "особенности функционирования",
    "источник и пояснение",
]


def normalize_rows(data: object) -> list[list[str]]:
    rows: list[list[str]] = []
    if not isinstance(data, list):
        return rows
    for item in data:
        if isinstance(item, dict):
            row = [
                " ".join(str(item.get("fragment", "")).split()),
                " ".join(str(item.get("biblicism", "")).split()),
                " ".join(str(item.get("classification", "")).split()),
                " ".join(str(item.get("features", "")).split()),
                " ".join(str(item.get("source_explanation", "")).split()),
            ]
        elif isinstance(item, list):
            row = [" ".join(str(cell).split()) for cell in item[:5]]
            while len(row) < 5:
                row.append("")
        else:
            continue
        if any(row):
            rows.append(row)
    return rows


def verify_docx(path: Path, template_path: Path) -> tuple[bool, str]:
    if not path.exists():
        return False, "output file was not created"
    if path.stat().st_size == 0:
        return False, "output file is empty"

    import docx

    try:
        doc = docx.Document(str(path))
        template = docx.Document(str(template_path))
    except Exception as exc:
        return False, f"docx open failed: {exc}"

    if not doc.tables or not template.tables:
        return False, "docx contains no tables"
    first_table = doc.tables[0]
    template_table = template.tables[0]
    if len(first_table.columns) != len(template_table.columns):
        return False, "docx table column count differs from template"
    if len(first_table.rows) < 2:
        return False, "docx table has no data rows"
    return True, "ok"


def write_json(rows: list[list[str]], path: Path) -> None:
    payload = [
        dict(
            zip(
                [
                    "fragment",
                    "biblicism",
                    "classification",
                    "features",
                    "source_explanation",
                ],
                row,
            )
        )
        for row in rows
    ]
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(rows: list[list[str]], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        handle.write("sep=;\n")
        writer = csv.writer(handle, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(HEADERS)
        writer.writerows(rows)


def write_md(rows: list[list[str]], path: Path) -> None:
    lines = [
        "| " + " | ".join(HEADERS) + " |",
        "| " + " | ".join([":---"] * len(HEADERS)) + " |",
    ]
    for row in rows:
        safe = [cell.replace("|", "\\|") for cell in row]
        lines.append("| " + " | ".join(safe) + " |")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_xlsx(rows: list[list[str]], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Workbook has no active sheet")
    sheet.title = "Biblicisms"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    header_font = Font(bold=True)
    widths = [28, 18, 24, 26, 28]

    for column_index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.border = border
        cell.alignment = wrap
        cell.font = header_font
        sheet.column_dimensions[cell.column_letter].width = widths[column_index - 1]

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.border = border
            cell.alignment = wrap

    workbook.save(path)


def run_word_maker(
    word_maker: Path, template: Path, output: Path, data_json: Path, source: Path
) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            sys.executable,
            "-X",
            "utf8",
            str(word_maker),
            str(template),
            str(output),
            str(data_json),
            str(source),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit(
            "Usage: generate_output.py <template.docx> <output.docx> <data.json> <source-file>"
        )

    template = Path(sys.argv[1])
    output = Path(sys.argv[2])
    data_json = Path(sys.argv[3])
    source = Path(sys.argv[4])
    output.parent.mkdir(parents=True, exist_ok=True)

    rows = normalize_rows(json.loads(data_json.read_text(encoding="utf-8")))
    base = output.with_suffix("")
    json_backup = base.with_suffix(".json")
    csv_backup = base.with_suffix(".csv")
    md_backup = base.with_suffix(".md")
    xlsx_backup = base.with_suffix(".xlsx")

    script_dir = Path(__file__).resolve().parent
    word_maker = script_dir / "word_maker.py"

    attempts: list[str] = []
    for _ in range(2):
        result = run_word_maker(word_maker, template, output, data_json, source)
        if result.returncode != 0:
            attempts.append((result.stdout + "\n" + result.stderr).strip())
        ok, message = verify_docx(output, template)
        if ok:
            write_json(rows, json_backup)
            write_csv(rows, csv_backup)
            write_md(rows, md_backup)
            write_xlsx(rows, xlsx_backup)
            print(f"OK: {output}")
            return
        attempts.append(message)

    write_json(rows, json_backup)
    write_csv(rows, csv_backup)
    write_md(rows, md_backup)
    write_xlsx(rows, xlsx_backup)

    print(f"FALLBACK_DATA: {json_backup} {csv_backup} {md_backup} {xlsx_backup}")
    print("DIAGNOSTIC: " + " | ".join(filter(None, attempts)))
    raise SystemExit(1)


if __name__ == "__main__":
    main()
