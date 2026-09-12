from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side


SUPPORTED_SOURCE_SUFFIXES = {".pdf", ".docx", ".xls", ".xlsx", ".txt"}
STATE_FILE = "pipeline_state.json"
HEADERS = [
    "фрагмент текста",
    "библеизм",
    "структурная классификация",
    "особенности функционирования",
    "источник и пояснение",
]
DICT_KEYS = [
    "fragment",
    "biblicism",
    "classification",
    "features",
    "source_explanation",
]

CHUNK_SIZES = {
    ".pdf": 25,
    ".docx": 250,
    ".xls": 300,
    ".xlsx": 300,
    ".txt": 300,
}

STRICT_CLASSIFICATION_MARKERS = (
    "антропоним",
    "топоним",
    "реали",
    "фразеолог",
    "цитат",
    "пословиц",
    "междомет",
    "phraseolog",
    "quotation",
    "proverb",
    "interjection",
)

CONTEXTUAL_CLASSIFICATION_MARKERS = (
    "фразеолог",
    "цитат",
    "пословиц",
    "аллюз",
    "парафраз",
    "афоризм",
    "phraseolog",
    "quotation",
    "proverb",
    "allusion",
    "paraphrase",
    "aphorism",
)

GENERIC_SYMBOL_WORDS = {
    "dragon",
    "snake",
    "star",
    "stars",
    "door",
    "garden",
    "tree",
    "fruit",
    "light",
    "darkness",
    "water",
    "fire",
    "road",
    "path",
    "gate",
    "serpent",
    "dragon",
    "star",
    "stars",
    "змея",
    "дракон",
    "звезда",
    "звезды",
    "звёзды",
    "дверь",
    "сад",
    "дерево",
    "плод",
    "свет",
    "тьма",
    "вода",
    "огонь",
    "путь",
    "ворота",
}

BARE_ARCHAISM_WORDS = {
    "thou",
    "thee",
    "thy",
    "thine",
    "th'art",
    "thart",
    "ye",
}

BIBLICAL_EVIDENCE_MARKERS = (
    "библ",
    "biblic",
    "евангел",
    "gospel",
    "ветхозав",
    "новозав",
    "genesis",
    "exodus",
    "psalm",
    "matthew",
    "mark",
    "luke",
    "john",
    "romans",
    "corinth",
    "исаия",
    "быт",
    "исх",
    "псал",
    "откров",
    "моис",
    "авраам",
    "христ",
    "иисус",
)

EXPLICIT_ALLUSION_MARKERS = (
    "аллюз",
    "отсыл",
    "парафраз",
    "цитат",
    "цитир",
    "фразеолог",
    "афоризм",
    "евангельск",
    "библейск",
    "ветхозав",
    "новозав",
    "allusion",
    "alludes",
    "reference to",
    "refers to",
    "paraphrase",
    "quote",
    "quotation",
    "biblical phrase",
    "biblical formula",
    "gospel",
)


def workspace_path(raw: str | None) -> Path:
    if raw:
        return Path(raw)
    return Path.cwd() / "biblicism_workspace"


def ensure_workspace(base: Path) -> None:
    for name in ("source", "temp", "example", "output", "context"):
        (base / name).mkdir(parents=True, exist_ok=True)


def save_state(base: Path, payload: dict[str, object]) -> None:
    (base / "temp" / STATE_FILE).write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def load_state(base: Path) -> dict[str, object]:
    path = base / "temp" / STATE_FILE
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def list_files(folder: Path, suffixes: set[str] | None = None) -> list[Path]:
    if not folder.exists():
        return []
    files = [p for p in folder.iterdir() if p.is_file()]
    if suffixes is not None:
        files = [p for p in files if p.suffix.lower() in suffixes]
    return sorted(files)


def first_file(folder: Path, suffixes: set[str] | None = None) -> Path | None:
    files = list_files(folder, suffixes)
    return files[0] if files else None


def run_python(script: Path, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-X", "utf8", str(script), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )


def normalize_text(value: object) -> str:
    return " ".join(str(value or "").split())


def int_from_state(value: object, default: int = 0) -> int:
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    if isinstance(value, str):
        try:
            return int(value)
        except ValueError:
            return default
    return default


def combined_item_text(item: dict[str, object]) -> str:
    return " ".join(
        normalize_text(item.get(key, ""))
        for key in DICT_KEYS
        if normalize_text(item.get(key, ""))
    ).lower()


def evidence_text(item: dict[str, object]) -> str:
    return " ".join(
        normalize_text(item.get(key, ""))
        for key in ("fragment", "biblicism", "features", "source_explanation")
        if normalize_text(item.get(key, ""))
    ).lower()


def has_biblical_evidence(item: dict[str, object]) -> bool:
    text = evidence_text(item)
    if any(marker in text for marker in BIBLICAL_EVIDENCE_MARKERS):
        return True
    return bool(
        re.search(
            r"\b(?:gen|exod|lev|num|deut|matt?|mark|luke|john|rom|cor|ps|rev)\.?\s*\d+[:.]\d+",
            text,
        )
        or re.search(
            r"\b(?:быт|исх|лев|чис|втор|мф|мк|лк|ин|рим|кор|пс|откр)\.?\s*\d+[:.]\d+",
            text,
        )
    )


def has_strict_classification(item: dict[str, object]) -> bool:
    classification = normalize_text(item.get("classification", "")).lower()
    return any(marker in classification for marker in STRICT_CLASSIFICATION_MARKERS)


def has_contextual_classification(item: dict[str, object]) -> bool:
    classification = normalize_text(item.get("classification", "")).lower()
    return any(marker in classification for marker in CONTEXTUAL_CLASSIFICATION_MARKERS)


def has_explicit_allusion_context(item: dict[str, object]) -> bool:
    text = " ".join(
        [
            normalize_text(item.get("classification", "")).lower(),
            normalize_text(item.get("features", "")).lower(),
            normalize_text(item.get("source_explanation", "")).lower(),
        ]
    )
    return any(
        marker in text for marker in EXPLICIT_ALLUSION_MARKERS
    ) or has_biblical_evidence(item)


def is_generic_symbol_only(item: dict[str, object]) -> bool:
    biblicism = normalize_text(item.get("biblicism", "")).lower()
    if not biblicism or biblicism not in GENERIC_SYMBOL_WORDS:
        return False

    context = " ".join(
        [
            normalize_text(item.get("classification", "")).lower(),
            normalize_text(item.get("features", "")).lower(),
            normalize_text(item.get("source_explanation", "")).lower(),
        ]
    )
    return "символ" in context or "symbol" in context


def is_bare_archaism(item: dict[str, object]) -> bool:
    biblicism = normalize_text(item.get("biblicism", "")).lower()
    if not biblicism or biblicism not in BARE_ARCHAISM_WORDS:
        return False
    return not has_explicit_allusion_context(item)


def is_valid_biblicism_item(item: object) -> bool:
    if not isinstance(item, dict):
        return False

    fragment = normalize_text(item.get("fragment", ""))
    biblicism = normalize_text(item.get("biblicism", ""))
    if not fragment or not biblicism:
        return False

    generic_symbol = is_generic_symbol_only(item)
    bare_archaism = is_bare_archaism(item)
    explicit_context = has_explicit_allusion_context(item)
    contextual_class = has_contextual_classification(item)
    biblical_evidence = has_biblical_evidence(item)

    if bare_archaism:
        return False

    if generic_symbol:
        return explicit_context and contextual_class and biblical_evidence

    if has_strict_classification(item):
        if not contextual_class and not biblical_evidence and not explicit_context:
            return False
        return True

    return biblical_evidence and explicit_context


def deduplicate_rows(data: list[dict[str, Any]]) -> list[dict[str, str]]:
    seen: dict[str, dict[str, str]] = {}
    for item in data:
        if not isinstance(item, dict):
            continue
        key = normalize_text(item.get("biblicism", "")).lower()
        if not key:
            continue
        if key in seen:
            existing_frag = seen[key].get("fragment", "")
            new_frag = normalize_text(item.get("fragment", ""))
            if new_frag and new_frag not in existing_frag:
                seen[key]["fragment"] = existing_frag + " | " + new_frag
        else:
            seen[key] = {k: normalize_text(item.get(k, "")) for k in DICT_KEYS}
    return list(seen.values())


def write_csv_output(rows: list[dict[str, str]], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        f.write("sep=;\n")
        writer = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(HEADERS)
        for row in rows:
            writer.writerow([row.get(k, "") for k in DICT_KEYS])


def write_xlsx_output(rows: list[dict[str, str]], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Workbook has no active sheet")
    sheet.title = "Biblicisms"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    header_font = Font(bold=True)
    widths = [28, 18, 24, 26, 28]

    for column_index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.border = border
        cell.alignment = wrap
        cell.font = header_font
        sheet.column_dimensions[cell.column_letter].width = widths[column_index - 1]

    for row_index, row in enumerate(rows, start=2):
        for column_index, key in enumerate(DICT_KEYS, start=1):
            cell = sheet.cell(
                row=row_index, column=column_index, value=row.get(key, "")
            )
            cell.border = border
            cell.alignment = wrap

    workbook.save(path)


def write_md_output(rows: list[dict[str, str]], path: Path) -> None:
    lines = [
        "| " + " | ".join(HEADERS) + " |",
        "| " + " | ".join([":---"] * len(HEADERS)) + " |",
    ]
    for row in rows:
        cells = [row.get(k, "").replace("|", "\\|") for k in DICT_KEYS]
        lines.append("| " + " | ".join(cells) + " |")
    path.write_text("\n".join(lines), encoding="utf-8")


def get_chunk_size(suffix: str) -> int:
    return CHUNK_SIZES.get(suffix.lower(), 25)


def get_total_units(script_dir: Path, source: Path) -> int:
    reader = script_dir / "read_pages.py"
    result = run_python(reader, "--file", str(source), "--start", "1", "--end", "1")
    for line in result.stdout.splitlines():
        if "Total pages" in line:
            return int(line.split("Total pages in PDF:")[1].split(".")[0].strip())
        if "Total paragraphs" in line:
            return int(line.split("Total paragraphs in DOCX:")[1].split(".")[0].strip())
        if "Total lines" in line:
            return int(line.split("Total lines in TXT:")[1].split(".")[0].strip())
        if "Total rows" in line:
            return int(line.split("Total rows in workbook:")[1].split(".")[0].strip())
    return 0


def command_init(args: argparse.Namespace) -> int:
    base = workspace_path(args.workspace)
    ensure_workspace(base)
    state = {
        "workspace": str(base),
        "source_files": [
            str(p) for p in list_files(base / "source", SUPPORTED_SOURCE_SUFFIXES)
        ],
        "example_files": [str(p) for p in list_files(base / "example")],
        "context_files": [str(p) for p in list_files(base / "context")],
        "chunks_processed": 0,
        "total_biblicisms": 0,
    }
    save_state(base, state)
    print(str(base))
    return 0


def command_status(args: argparse.Namespace) -> int:
    base = workspace_path(args.workspace)
    ensure_workspace(base)
    state = load_state(base)

    data_json = base / "temp" / "data.json"
    total_found = 0
    if data_json.exists():
        try:
            total_found = len(json.loads(data_json.read_text(encoding="utf-8")))
        except Exception:
            pass

    status = {
        "workspace": str(base),
        "source_files": [
            str(p) for p in list_files(base / "source", SUPPORTED_SOURCE_SUFFIXES)
        ],
        "example_files": [str(p) for p in list_files(base / "example")],
        "context_files": [str(p) for p in list_files(base / "context")],
        "output_files": [str(p) for p in list_files(base / "output")],
        "chunks_processed": state.get("chunks_processed", 0),
        "total_biblicisms_found": total_found,
    }
    print(json.dumps(status, ensure_ascii=False, indent=2))
    return 0


def command_read(args: argparse.Namespace) -> int:
    base = workspace_path(args.workspace)
    ensure_workspace(base)
    script_dir = Path(__file__).resolve().parent
    reader = script_dir / "read_pages.py"

    target_folder = base / args.target
    source = (
        Path(args.file)
        if args.file
        else first_file(
            target_folder,
            None if args.target == "context" else SUPPORTED_SOURCE_SUFFIXES,
        )
    )

    if source is None or not source.exists():
        print(f"ERROR: no readable file found in {target_folder}", file=sys.stderr)
        return 1

    output_name = f"{source.stem}_{args.start}_{args.end}.txt"
    output_path = base / "temp" / output_name

    result = run_python(
        reader,
        "--file",
        str(source),
        "--start",
        str(args.start),
        "--end",
        str(args.end),
        "--out",
        str(output_path),
    )
    if result.returncode != 0:
        print(result.stdout)
        print(result.stderr, file=sys.stderr)
        return result.returncode

    state = load_state(base)
    state["last_read"] = {
        "source": str(source),
        "target": args.target,
        "start": args.start,
        "end": args.end,
        "artifact": str(output_path),
    }
    save_state(base, state)
    print(str(output_path))
    return 0


def command_append(args: argparse.Namespace) -> int:
    base = workspace_path(args.workspace)
    ensure_workspace(base)

    json_input = Path(args.json)
    if not json_input.exists():
        print(f"ERROR: JSON file not found: {json_input}", file=sys.stderr)
        return 1

    new_data = json.loads(json_input.read_text(encoding="utf-8"))
    if not isinstance(new_data, list):
        print("ERROR: JSON must be a list of objects", file=sys.stderr)
        return 1

    filtered_data = [item for item in new_data if is_valid_biblicism_item(item)]
    rejected_data = [item for item in new_data if not is_valid_biblicism_item(item)]

    data_json = base / "temp" / "data.json"
    existing: list[dict[str, Any]] = []
    if data_json.exists():
        try:
            loaded_existing = json.loads(data_json.read_text(encoding="utf-8"))
            if isinstance(loaded_existing, list):
                existing = [item for item in loaded_existing if isinstance(item, dict)]
        except Exception:
            existing = []

    merged = existing + filtered_data
    deduped = deduplicate_rows(merged)
    data_json.write_text(
        json.dumps(deduped, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    write_csv_output(deduped, base / "output" / "biblicisms.csv")
    write_xlsx_output(deduped, base / "output" / "biblicisms.xlsx")
    write_md_output(deduped, base / "output" / "biblicisms.md")

    rejected_path = base / "temp" / "rejected_last_append.json"
    rejected_path.write_text(
        json.dumps(rejected_data, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    state = load_state(base)
    chunks_done = int_from_state(state.get("chunks_processed", 0)) + 1
    state["chunks_processed"] = chunks_done
    state["total_biblicisms"] = len(deduped)
    save_state(base, state)

    total_chunks = state.get("total_chunks", "?")
    print(
        f"APPENDED: +{len(filtered_data)} accepted | {len(deduped)} unique total | chunk {chunks_done}/{total_chunks}"
    )
    print(f"REJECTED: {len(rejected_data)} weak/generic candidates")
    print(f"CSV: {base / 'output' / 'biblicisms.csv'}")
    print(f"XLSX: {base / 'output' / 'biblicisms.xlsx'}")
    print(f"MD:  {base / 'output' / 'biblicisms.md'}")
    print(f"REJECTS: {rejected_path}")
    return 0


def command_finalize(args: argparse.Namespace) -> int:
    base = workspace_path(args.workspace)
    ensure_workspace(base)
    script_dir = Path(__file__).resolve().parent
    generator = script_dir / "generate_output.py"

    template = (
        Path(args.template)
        if args.template
        else first_file(base / "example", {".docx"})
    )
    source = (
        Path(args.source)
        if args.source
        else first_file(base / "source", SUPPORTED_SOURCE_SUFFIXES)
    )
    data_json = Path(args.data) if args.data else base / "temp" / "data.json"

    if template is None or not template.exists():
        print("ERROR: template .docx not found", file=sys.stderr)
        return 1
    if source is None or not source.exists():
        print("ERROR: source file not found", file=sys.stderr)
        return 1
    if not data_json.exists():
        print("ERROR: data.json not found", file=sys.stderr)
        return 1

    output = (
        Path(args.output) if args.output else base / "output" / f"{source.stem}.docx"
    )
    result = run_python(
        generator, str(template), str(output), str(data_json), str(source)
    )
    print(result.stdout.strip())
    if result.returncode != 0:
        print(result.stderr, file=sys.stderr)
        return result.returncode

    state = load_state(base)
    state["last_output"] = str(output)
    save_state(base, state)
    return 0


def command_auto_info(args: argparse.Namespace) -> int:
    base = workspace_path(args.workspace)
    ensure_workspace(base)
    script_dir = Path(__file__).resolve().parent

    source = first_file(base / "source", SUPPORTED_SOURCE_SUFFIXES)
    if source is None:
        print("ERROR: no source file found", file=sys.stderr)
        return 1

    total = get_total_units(script_dir, source)
    chunk_size = get_chunk_size(source.suffix)

    chunks = []
    start = 1
    while start <= total:
        end = min(start + chunk_size - 1, total)
        chunks.append({"start": start, "end": end})
        start = end + 1

    info = {
        "source": str(source),
        "total_units": total,
        "chunk_size": chunk_size,
        "total_chunks": len(chunks),
        "chunks": chunks,
    }

    state = load_state(base)
    state["total_chunks"] = len(chunks)
    save_state(base, state)

    print(json.dumps(info, ensure_ascii=False, indent=2))
    return 0


def command_compare(args: argparse.Namespace) -> int:
    base = workspace_path(args.workspace)

    file_a = Path(args.run_a)
    file_b = Path(args.run_b)
    if not file_a.exists() or not file_b.exists():
        print("ERROR: one or both run files not found", file=sys.stderr)
        return 1

    data_a = json.loads(file_a.read_text(encoding="utf-8"))
    data_b = json.loads(file_b.read_text(encoding="utf-8"))

    keys_a = {
        normalize_text(item.get("biblicism", "")).lower()
        for item in data_a
        if isinstance(item, dict)
    }
    keys_b = {
        normalize_text(item.get("biblicism", "")).lower()
        for item in data_b
        if isinstance(item, dict)
    }

    only_a = sorted(keys_a - keys_b)
    only_b = sorted(keys_b - keys_a)
    common = sorted(keys_a & keys_b)

    result = {
        "run_a": str(file_a),
        "run_b": str(file_b),
        "common": len(common),
        "only_in_a": only_a,
        "only_in_b": only_b,
        "total_a": len(keys_a),
        "total_b": len(keys_b),
    }
    out = base / "output" / "comparison.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


# ─── Parser ─────────────────────────────────────────────────────


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Biblicism Extractor Pipeline")
    subparsers = parser.add_subparsers(dest="command", required=True)

    p = subparsers.add_parser("init", help="Create workspace folders")
    p.add_argument("--workspace")
    p.set_defaults(func=command_init)

    p = subparsers.add_parser("status", help="Show workspace state")
    p.add_argument("--workspace")
    p.set_defaults(func=command_status)

    p = subparsers.add_parser("read", help="Read a chunk from source or context")
    p.add_argument("--workspace")
    p.add_argument("--target", choices=["source", "context"], default="source")
    p.add_argument("--file")
    p.add_argument("--start", type=int, required=True)
    p.add_argument("--end", type=int, required=True)
    p.set_defaults(func=command_read)

    p = subparsers.add_parser("append", help="Append biblicisms JSON and export CSV/MD")
    p.add_argument("--workspace")
    p.add_argument("--json", required=True)
    p.set_defaults(func=command_append)

    p = subparsers.add_parser("finalize", help="Build DOCX from data.json + template")
    p.add_argument("--workspace")
    p.add_argument("--template")
    p.add_argument("--source")
    p.add_argument("--data")
    p.add_argument("--output")
    p.set_defaults(func=command_finalize)

    p = subparsers.add_parser("auto", help="Get chunk plan for full-book processing")
    p.add_argument("--workspace")
    p.set_defaults(func=command_auto_info)

    p = subparsers.add_parser("compare", help="Compare two analysis runs")
    p.add_argument("--workspace")
    p.add_argument("--run-a", required=True)
    p.add_argument("--run-b", required=True)
    p.set_defaults(func=command_compare)

    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()
    raise SystemExit(args.func(args))


if __name__ == "__main__":
    main()
